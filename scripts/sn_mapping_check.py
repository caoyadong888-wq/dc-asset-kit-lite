#!/usr/bin/env python3
"""Compare an asset register with collected serial dumps.

Flags MISSING, UNKNOWN, DUPLICATE, MISMATCH_LOCATION, and MISMATCH_PN.
Exit 0 when the register matches the dumps. Exit 2 when there are findings.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

from sn_dump_parse import (
    DumpRecord,
    canonicalize_port,
    norm_host,
    norm_pn,
    norm_sn,
    parse_dump,
    short_source,
    text_table,
)


FINDING_TYPES = (
    "MISSING",
    "UNKNOWN",
    "DUPLICATE",
    "MISMATCH_LOCATION",
    "MISMATCH_PN",
)

SEVERITY = {
    "MISSING": "high",
    "UNKNOWN": "medium",
    "DUPLICATE": "high",
    "MISMATCH_LOCATION": "high",
    "MISMATCH_PN": "medium",
}


@dataclass
class RegisterRecord:
    asset_id: str
    sn: str
    pn: str
    host: str
    port: str
    rack: str = ""
    u_position: str = ""
    status: str = ""
    asset_class: str = "optic"
    source: str = ""
    vendor: str = ""


@dataclass
class Finding:
    type: str
    severity: str
    sn: str
    asset_ids: list[str] = field(default_factory=list)
    asset_class: str = ""
    register_pn: str = ""
    dump_pn: str = ""
    register_host: str = ""
    register_port: str = ""
    dump_host: str = ""
    dump_port: str = ""
    sources: list[str] = field(default_factory=list)
    detail: str = ""


def expand_inputs(items: list[str], suffixes: tuple[str, ...]) -> list[Path]:
    found: list[Path] = []
    allowed = tuple(suffix.lower() for suffix in suffixes)
    for item in items:
        path = Path(item)
        if path.is_dir():
            for child in sorted(path.rglob("*")):
                if not child.is_file():
                    continue
                if child.name.startswith("~$"):
                    continue
                if child.suffix.lower() in allowed:
                    found.append(child)
        elif path.is_file():
            found.append(path)
        else:
            raise SystemExit(f"not found: {item}")
    return found


def _norm_header(value: str) -> str:
    text = (value or "").strip().lower().replace("-", "_").replace(" ", "_")
    return text


def infer_asset_class(path: Path, sheet: str, headers: set[str]) -> str:
    blob = f"{path.name} {sheet}".lower()
    if "switch" in blob:
        return "switch"
    if "aoc" in blob:
        return "aoc"
    if "optic" in blob:
        return "optic"
    if "hostname" in headers or "u_position" in headers:
        return "switch"
    return "optic"


def _row_get(row: dict[str, str], *names: str) -> str:
    for name in names:
        if name in row and row[name] is not None:
            return str(row[name]).strip()
    return ""


def records_from_rows(
    rows: list[dict[str, str]],
    path: Path,
    sheet: str = "",
) -> list[RegisterRecord]:
    if not rows:
        return []
    headers = set(rows[0].keys())
    asset_class = infer_asset_class(path, sheet, headers)
    records: list[RegisterRecord] = []
    for row in rows:
        raw_sn = _row_get(row, "sn", "serial")
        sn = norm_sn(raw_sn)
        if not sn:
            continue
        host = norm_host(_row_get(row, "hostname", "host"))
        pn = norm_pn(_row_get(row, "pn", "model"))
        port = ""
        if asset_class != "switch":
            port = canonicalize_port(_row_get(row, "port"))
        asset_id = _row_get(row, "asset_id") or sn
        explicit = _row_get(row, "asset_class")
        records.append(
            RegisterRecord(
                asset_id=asset_id,
                sn=sn,
                pn=pn,
                host=host,
                port=port,
                rack=_row_get(row, "rack"),
                u_position=_row_get(row, "u_position"),
                status=_row_get(row, "status").lower(),
                asset_class=(explicit or asset_class).lower(),
                source=str(path),
                vendor=_row_get(row, "vendor"),
            )
        )
    return records


def _read_csv_dicts(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows: list[dict[str, str]] = []
        for raw in reader:
            item: dict[str, str] = {}
            for key, value in raw.items():
                if key is None:
                    continue
                item[_norm_header(key)] = "" if value is None else str(value)
            if any(cell.strip() for cell in item.values()):
                rows.append(item)
        return rows


def load_register_file(path: Path) -> list[RegisterRecord]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return records_from_rows(_read_csv_dicts(path), path)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        records: list[RegisterRecord] = []
        try:
            for sheet in workbook.worksheets:
                iterator = sheet.iter_rows(values_only=True)
                try:
                    header_row = next(iterator)
                except StopIteration:
                    continue
                headers = [_norm_header(str(cell)) if cell is not None else "" for cell in header_row]
                while headers and headers[-1] == "":
                    headers.pop()
                dict_rows: list[dict[str, str]] = []
                for values in iterator:
                    item = {}
                    blank = True
                    for index, header in enumerate(headers):
                        if not header:
                            continue
                        cell = values[index] if index < len(values) else ""
                        text = "" if cell is None else str(cell)
                        if text.strip():
                            blank = False
                        item[header] = text
                    if not blank:
                        dict_rows.append(item)
                records.extend(records_from_rows(dict_rows, path, sheet.title))
        finally:
            workbook.close()
        return records
    raise SystemExit(f"unsupported register file: {path}")


def load_register(items: list[str]) -> list[RegisterRecord]:
    records: list[RegisterRecord] = []
    for path in expand_inputs(items, (".csv", ".xlsx")):
        records.extend(load_register_file(path))
    return records


def load_dumps(items: list[str]) -> list[DumpRecord]:
    records: list[DumpRecord] = []
    for path in expand_inputs(items, (".txt", ".json", ".log")):
        records.extend(parse_dump(path))
    return records


def _hosts_agree(left: str, right: str) -> bool:
    return left.strip().casefold() == right.strip().casefold()


def _ports_agree(left: str, right: str) -> bool:
    return left.strip().casefold() == right.strip().casefold()


def _location_match(reg: RegisterRecord, dump: DumpRecord) -> bool:
    return _hosts_agree(reg.host, dump.host) and _ports_agree(reg.port, dump.port)


def _pn_match(reg: RegisterRecord, dump: DumpRecord) -> bool:
    return norm_pn(reg.pn) == norm_pn(dump.pn)


def _fields(pairs: list[tuple[str, str]]) -> str:
    return " ".join(f"{key}={value}" for key, value in pairs if str(value).strip())


def _join_unique(values: list[str]) -> str:
    seen: list[str] = []
    for value in values:
        if value not in seen:
            seen.append(value)
    return " | ".join(seen)


def check_mapping(register: list[RegisterRecord], dumps: list[DumpRecord]) -> list[Finding]:
    by_sn: dict[str, list[RegisterRecord]] = {}
    for record in register:
        by_sn.setdefault(record.sn, []).append(record)
    dump_by_sn: dict[str, list[DumpRecord]] = {}
    for record in dumps:
        if record.sn:
            dump_by_sn.setdefault(record.sn, []).append(record)

    findings: list[Finding] = []
    for sn, rows in by_sn.items():
        asset_ids = sorted({row.asset_id for row in rows})
        classes = _join_unique([row.asset_class for row in rows])
        sources = [short_source(row.source) for row in rows]
        seen = dump_by_sn.get(sn, [])
        if len(rows) > 1:
            findings.append(
                Finding(
                    type="DUPLICATE",
                    severity=SEVERITY["DUPLICATE"],
                    sn=sn,
                    asset_ids=asset_ids,
                    asset_class=classes,
                    register_pn=_join_unique([row.pn for row in rows]),
                    register_host=_join_unique([row.host for row in rows]),
                    register_port=_join_unique([row.port for row in rows]),
                    sources=sources,
                    detail=f"sn repeated in register; asset_ids={'|'.join(asset_ids)}",
                )
            )
        if not seen:
            findings.append(
                Finding(
                    type="MISSING",
                    severity=SEVERITY["MISSING"],
                    sn=sn,
                    asset_ids=asset_ids,
                    asset_class=classes,
                    register_pn=_join_unique([row.pn for row in rows]),
                    register_host=_join_unique([row.host for row in rows]),
                    register_port=_join_unique([row.port for row in rows]),
                    sources=sources,
                    detail=(
                        "in register, not in dumps; "
                        f"asset_ids={'|'.join(asset_ids)}; "
                        + _fields([
                            ("host", rows[0].host),
                            ("port", rows[0].port),
                            ("pn", rows[0].pn),
                        ])
                    ),
                )
            )
            continue
        if not any(_location_match(row, dump) for row in rows for dump in seen):
            reg_locs = _join_unique([f"{row.host} {row.port}".strip() for row in rows])
            dump_locs = _join_unique([f"{dump.host} {dump.port}".strip() for dump in seen])
            findings.append(
                Finding(
                    type="MISMATCH_LOCATION",
                    severity=SEVERITY["MISMATCH_LOCATION"],
                    sn=sn,
                    asset_ids=asset_ids,
                    asset_class=classes,
                    register_pn=_join_unique([row.pn for row in rows]),
                    dump_pn=_join_unique([dump.pn for dump in seen]),
                    register_host=_join_unique([row.host for row in rows]),
                    register_port=_join_unique([row.port for row in rows]),
                    dump_host=_join_unique([dump.host for dump in seen]),
                    dump_port=_join_unique([dump.port for dump in seen]),
                    sources=[short_source(dump.source) for dump in seen],
                    detail=f"host/port differ; register {reg_locs}; dump {dump_locs}",
                )
            )
        if not any(_pn_match(row, dump) for row in rows for dump in seen):
            reg_pn = _join_unique([row.pn for row in rows])
            dump_pn = _join_unique([dump.pn for dump in seen])
            findings.append(
                Finding(
                    type="MISMATCH_PN",
                    severity=SEVERITY["MISMATCH_PN"],
                    sn=sn,
                    asset_ids=asset_ids,
                    asset_class=classes,
                    register_pn=reg_pn,
                    dump_pn=dump_pn,
                    register_host=_join_unique([row.host for row in rows]),
                    register_port=_join_unique([row.port for row in rows]),
                    dump_host=_join_unique([dump.host for dump in seen]),
                    dump_port=_join_unique([dump.port for dump in seen]),
                    sources=[short_source(dump.source) for dump in seen],
                    detail=f"pn/model differ; register {reg_pn}; dump {dump_pn}",
                )
            )

    for sn, seen in dump_by_sn.items():
        if sn in by_sn:
            continue
        first = seen[0]
        findings.append(
            Finding(
                type="UNKNOWN",
                severity=SEVERITY["UNKNOWN"],
                sn=sn,
                dump_pn=_join_unique([dump.pn for dump in seen]),
                dump_host=_join_unique([dump.host for dump in seen]),
                dump_port=_join_unique([dump.port for dump in seen]),
                sources=[short_source(dump.source) for dump in seen],
                detail=(
                    "seen in dumps, not in register; "
                    + _fields([
                        ("host", first.host),
                        ("port", first.port),
                        ("pn", first.pn),
                        ("source", short_source(first.source)),
                    ])
                ),
            )
        )

    order = {name: index for index, name in enumerate(FINDING_TYPES)}
    findings.sort(key=lambda item: (order[item.type], item.sn, "|".join(item.asset_ids)))
    return findings


def count_findings(findings: list[Finding]) -> dict[str, int]:
    counts = {name: 0 for name in FINDING_TYPES}
    for finding in findings:
        counts[finding.type] += 1
    return counts


def exit_code_for(findings: list[Finding]) -> int:
    return 2 if findings else 0


def _rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return str(path)


def summary_lines(
    findings: list[Finding],
    register: list[RegisterRecord],
    dumps: list[DumpRecord],
) -> str:
    counts = count_findings(findings)
    unique_reg = len({row.sn for row in register})
    unique_dump = len({row.sn for row in dumps})
    code = exit_code_for(findings)
    head = [
        "== SN mapping check ==",
        f"register_rows: {len(register)}",
        f"unique_register_sn: {unique_reg}",
        f"dump_records: {len(dumps)}",
        f"unique_dump_sn: {unique_dump}",
        f"findings: {len(findings)}",
        f"exit: {code}",
        "",
        text_table(
            ["finding", "count"],
            [[name, counts[name]] for name in FINDING_TYPES],
        ),
        "",
        text_table(
            ["type", "sn", "asset_ids", "detail"],
            [
                [item.type, item.sn, "|".join(item.asset_ids), item.detail]
                for item in findings
            ]
            or [["(none)", "", "", "register matches dumps"]],
        ),
    ]
    return "\n".join(head)


def write_reports(
    findings: list[Finding],
    register: list[RegisterRecord],
    dumps: list[DumpRecord],
    out_dir: Path,
) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    counts = count_findings(findings)
    md_path = out_dir / "report.md"
    csv_path = out_dir / "report.csv"
    json_path = out_dir / "report.json"

    md_lines = [
        "# SN mapping report",
        "",
        summary_lines(findings, register, dumps),
        "",
        "SN comparison is strip + uppercase. Host and port comparison is case-insensitive.",
        "Switch rows compare hostname only. AOC and optic rows compare host and port.",
        "A location or part-number mismatch is raised when no register row for that SN agrees with a dump.",
        "",
    ]
    for name in FINDING_TYPES:
        group = [item for item in findings if item.type == name]
        md_lines.append(f"## {name} ({len(group)})")
        md_lines.append("")
        if not group:
            md_lines.append("None.")
            md_lines.append("")
            continue
        for item in group:
            md_lines.append(f"- `{item.sn}` — {item.detail}")
        md_lines.append("")
    md_path.write_text("\n".join(md_lines), encoding="utf-8")

    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        fieldnames = [
            "type",
            "severity",
            "sn",
            "asset_ids",
            "asset_class",
            "register_pn",
            "dump_pn",
            "register_host",
            "register_port",
            "dump_host",
            "dump_port",
            "sources",
            "detail",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for item in findings:
            writer.writerow(
                {
                    "type": item.type,
                    "severity": item.severity,
                    "sn": item.sn,
                    "asset_ids": "|".join(item.asset_ids),
                    "asset_class": item.asset_class,
                    "register_pn": item.register_pn,
                    "dump_pn": item.dump_pn,
                    "register_host": item.register_host,
                    "register_port": item.register_port,
                    "dump_host": item.dump_host,
                    "dump_port": item.dump_port,
                    "sources": "|".join(item.sources),
                    "detail": item.detail,
                }
            )

    payload = {
        "register_rows": len(register),
        "unique_register_sn": len({row.sn for row in register}),
        "dump_records": len(dumps),
        "unique_dump_sn": len({row.sn for row in dumps}),
        "findings_total": len(findings),
        "counts": counts,
        "exit_code": exit_code_for(findings),
        "findings": [asdict(item) for item in findings],
    }
    json_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return [md_path, csv_path, json_path]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Check register serials against collected dumps.")
    parser.add_argument("--register", nargs="+", required=True, help="CSV and/or XLSX register files or a directory")
    parser.add_argument("--dumps", nargs="+", required=True, help="Dump files or a directory of txt/json dumps")
    parser.add_argument("--out", required=True, help="Directory for report.md, report.csv, and report.json")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        register = load_register(args.register)
        dumps = load_dumps(args.dumps)
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    findings = check_mapping(register, dumps)
    written = write_reports(findings, register, dumps, Path(args.out))
    print(summary_lines(findings, register, dumps))
    print()
    for path in written:
        print(f"wrote {_rel(path)}")
    return exit_code_for(findings)


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    raise SystemExit(main())
